"""Build the course presentation from checked runtime evidence (no secrets)."""
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]


def read(name):
    path = ROOT / 'reports' / name
    return json.loads(path.read_text(encoding='utf-8')) if path.exists() else {}


def build():
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    navy, white, teal = RGBColor(14,28,48), RGBColor(242,246,252), RGBColor(61,210,182)
    coverage = read('coverage-maintenance.json').get('totals',{}).get('percent_covered',0)
    monitoring, verification, recovery = (read(n) for n in ('monitoring-verification.json','verification.json','recovery-verification.json'))
    actions = read('github-actions.json').get('runs',[])

    def text(slide, x, y, w, h, content, size=24, color=white):
        box = slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        box.text_frame.word_wrap = True
        for i,line in enumerate(content.split('\n')):
            p = box.text_frame.paragraphs[0] if i == 0 else box.text_frame.add_paragraph()
            p.text = line
            p.font.name, p.font.size, p.font.color.rgb = 'Arial', Pt(size), color
            p.space_after = Pt(14)
        return box

    def slide(title, lines, notes):
        s = deck.slides.add_slide(deck.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = navy
        text(s,.65,.45,12,1,title,34,teal)
        text(s,.85,1.65,11.7,4.95,lines,24)
        text(s,.65,7.03,12,.3,f'DDM501  ·  Face Voice Proctoring                                    {len(deck.slides):02}',11)
        s.notes_slide.notes_text_frame.text = notes
        return s

    slide('Face + Voice Proctoring',
          'API xác minh nhân viên trong kỳ đánh giá ngoại ngữ\nMLOps: từ dữ liệu tới serving, monitoring và vận hành\nDemo Docker Desktop · Grafana · Telegram · GitHub Actions',
          '1 phút. Điền tên thật của nhóm trên slide nếu có. Giới thiệu batch API và portal công ty. Capability limits rõ ràng; không công bố production accuracy.')
    slide('Vấn đề, người dùng và tiêu chí thành công',
          'Công ty có phần mềm thi và chấm điểm riêng.\nDịch vụ nhận ảnh/WAV theo nhịp do khách hàng cấu hình.\nAPI/webhook trả identity và integrity signals.\nAdmin công ty xem lịch sử, bằng chứng và PDF/CSV.',
          '1,5 phút. Chỉ tiêu kinh doanh là đề xuất pilot, chưa đo ở khách hàng. Gate 20% là demo và phải được thay bằng chính sách chấp nhận rủi ro khi có dữ liệu thật.')
    s = slide('Full pipeline và các ranh giới trách nhiệm','',
              '2 phút. Mỗi block liên kết với code/DAG/report. Enrollment tạo embeddings; Airflow orchestration dùng snapshot đã fingerprint; Registry điều khiển threshold. Grafana gom telemetry, không thay chức năng điều phối hoặc experiment tracking.')
    for i,(title,detail) in enumerate([('Data','Consent + enrollment'),('Pipeline','Snapshot → validate'),('Experiments','Identity CV + holdout'),('Registry','RAI → gate → champion'),('Serving','Batch API + evidence'),('Monitoring','Grafana → Telegram')]):
        x,y = .8+(i%3)*4.2, 1.8+(i//3)*2.3
        shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(3.8),Inches(1.6))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(26,53,77)
        shape.line.color.rgb = teal
        text(s,x+.2,y+.15,3.4,.45,f'{i+1}. {title}',23,teal)
        text(s,x+.2,y+.8,3.4,.6,detail,19)
    slide('Dữ liệu, lineage và chất lượng',
          'Snapshot bất biến theo DAG run; fingerprint SHA-256.\nTraining scope mặc định tenant demo.\nReject: invalid/zero vector, inconsistent dimension, duplicates.\nMLflow lưu snapshot, validation, thresholds và evaluation folds.',
          '1 phút. Không đưa tenant khác vào training mặc định. Bộ bootstrap LFW/Speech Commands ghép tổng hợp, không phải một tập face+voice thật có consent từ cùng người.')
    slide('Calibration và lựa chọn model',
          'Max-template cosine giống cách serving so sánh.\nFive-way identity split: holdout riêng + 4 folds CV nội bộ.\n1.151 thresholds × 3 objectives × 2 modalities.\nGate kiểm tra calibration, CV, holdout FAR/FRR và sample counts.',
          '1,5 phút. Mở evaluation/identity-disjoint.json và nested MLflow runs. CV bên trong phần calibration, fold 0 giữ ngoài tune. Đổi alias champion chỉ sau RAI audit và gate; không nới gate để có DAG xanh.')
    slide('Batch serving và portal doanh nghiệp',
          'Tenant-scoped API keys; role operator / integration / platform.\nBatch có consent, request_id và lịch sử bất biến.\nCallback HMAC + outbox + retry + dedup.\nDemo hai công ty: identity đúng/sai, evidence và export.',
          '2,5 phút. Mở customer demo :18600 và portal :18501. Ghi danh, gọi batch check, kiểm tra webhook và bằng chứng, xuất PDF/CSV. Quyết định thi thuộc khách hàng. Browser camera/microphone cần người dùng chấp nhận quyền và test trên thiết bị thật.')
    slide('Grafana là trung tâm monitoring',
          'Service readiness · p50/p95 · errors · training quality\nPSI/Evidently · Registry/CV/holdout · human/synthetic performance\nBatch/detector/evidence · webhooks · Airflow/RAI · CPU/RAM/logs\nReport cùng origin, yêu cầu đăng nhập Grafana.',
          '2 phút. Mở /d/biometric-overview. Dùng tenant/project/service selector. Grafana là trang quan sát chính; Airflow/MLflow giữ chức năng điều khiển. Human và synthetic là hai nguồn riêng.')
    slide('Cảnh báo và vận hành',
          'Prometheus rules → Alertmanager → ops-monitor → Telegram.\nFiring và resolved được gửi; lỗi transport dẫn tới retry.\nDrift >0,2; service/collector down; data invalid; webhook failed.\nBackup restore vào DB riêng; rollback alias rồi khôi phục champion.',
          '1 phút. Hiển thị bot ddm501_face_voice_proctoring_bot và bằng chứng delivery; không mở token. Restore drill không ghi đè DB thật. Rollback rehearsal gây đổi model ngắn trên demo rồi phục hồi model ban đầu.')
    slide('Responsible AI và explainability',
          'Reason codes + score margins + threshold sensitivity.\nCounterfactual thay một score, giữ quality và modality còn lại.\nHuman fairness cần ≥20 samples/slice và ≥5 mỗi class.\nWilson 95% CI; insufficient_data khi chưa có đủ nhãn.',
          '1,5 phút. Quality slice là proxy vận hành, không chứng minh demographic fairness. Policy explanation không phải lời giải thích causal cho embedding. PAD/AASIST chưa có benchmark khách hàng; media đã enroll có score cao không chứng minh khả năng generalize.')
    latest = actions[0] if actions else {}
    slide('Bằng chứng kiểm chứng',
          f'Coverage phạm vi API + monitoring + evaluation gates: {coverage:.2f}%.\nFull pipeline verification: {verification.get("status","chưa chạy")}.\nGrafana queries / reports / alert delivery: {monitoring.get("status","chưa chạy")}.\nRecovery drill: {recovery.get("backup_restore",{}).get("status","chưa chạy")}.\nGitHub Actions gần nhất: {latest.get("status","chưa chạy")} / {latest.get("conclusion","pending")}.',
          '1 phút. Dẫn đến reports JSON local và Actions URL. Coverage không bao gồm tất cả scripts CLI, frontend, weights hoặc Airflow; live integration bổ sung evidence. Nếu pending/failure, báo đúng trạng thái, không suy ra thành công từ workflow config.')
    slide('Scalability, cost và các mục cần nhóm hoàn tất',
          'Sizing: workers ≥ ceil(λ × CPU service time / 0,65).\nChi phí: compute + DB + suspicious evidence + egress.\nCần người thật: consented labels, camera/mic acceptance, team contribution.\nPAD/AASIST có inference; cần benchmark và cloud/TLS/SSO.',
          '1 phút. Sizing là mô hình để load test, không phải capacity đã chứng minh. Không dựng tên hay commit thành viên. Không có deployment cloud có phí trong phạm vi này.')
    slide('Q&A / demo links',
          'Grafana: localhost:13000/d/biometric-overview\nPortal: :18501 · Legacy: :18600 · API docs: :18100/docs\nAirflow: :18081 · MLflow: :15030\nGitHub: TrinhDucDuong/ddm501-face-voice-proctoring',
          'Q&A: Vì sao threshold calibration thay fine-tune? Vì base encoders pretrained và dữ liệu demo nhỏ. Vì sao identity split? Tránh cùng identity xuất hiện ở tune và test. Mất callback? Transactional outbox + retries + idempotency. Fairness? Insufficient human labels, không giả định đạt. Replay protocol khác spoof biometric. Chỉ partner backend có thể enforce quyền vào thi.')
    destination = ROOT / 'docs/DDM501_Face_Voice_Proctoring.pptx'
    destination.parent.mkdir(exist_ok=True)
    deck.save(destination)
    print(f'Created {len(deck.slides)} slides with speaker notes')


if __name__ == '__main__':
    build()
